"""Windows slide-rendering fallback tests (no course material required)."""
import os

import pytest
from pptx import Presentation

import ingest


@pytest.mark.skipif(os.name != "nt", reason="PowerPoint COM is Windows-only")
def test_powerpoint_unavailable_falls_back_to_text_only(monkeypatch, tmp_path):
    import win32com.client

    deck_path = tmp_path / "synthetic.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 3000000, 500000).text = "Hello slide"
    deck.save(deck_path)

    monkeypatch.setattr(ingest, "find_soffice", lambda: None)

    def unavailable(*args, **kwargs):
        raise RuntimeError("PowerPoint cannot start")

    monkeypatch.setattr(win32com.client, "DispatchEx", unavailable)
    slides, warning = ingest.render_pptx_slides(str(deck_path), str(tmp_path))
    assert len(slides) == 1
    assert slides[0]["text"] == "Hello slide"
    assert slides[0]["image_path"] is None
    assert warning and "text only" in warning.lower()
